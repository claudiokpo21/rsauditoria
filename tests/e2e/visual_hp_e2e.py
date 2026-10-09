"""
Prueba visual de la planilla de resultado H&P (colores de bandas) en pantalla, PDF y Excel,
con la plantilla importada del Excel y las respuestas de la planilla (6,27 – Bueno).
Chromium con escritorio y celular emulado; el servidor Supabase se simula.

Uso: python3 tests/e2e/visual_hp_e2e.py http://localhost:4177 [carpeta_salida]
"""
import json, os, sys, uuid
from playwright.sync_api import sync_playwright
from pypdf import PdfReader
from openpyxl import load_workbook
from fake_supabase import FakeSupabase, now_iso

BASE = sys.argv[1] if len(sys.argv) > 1 else 'http://localhost:4177'
OUT = sys.argv[2] if len(sys.argv) > 2 else 'e2e-out-visual'
os.makedirs(OUT, exist_ok=True)
SB = 'https://hhwfhearafhmougtfssu.supabase.co'
STORAGE_KEY = 'sb-hhwfhearafhmougtfssu-auth-token'
HERE = os.path.dirname(os.path.abspath(__file__))
results = []

def check(name, cond, detail=''):
    results.append({'prueba': name, 'ok': bool(cond), 'detalle': detail})
    print(('OK   ' if cond else 'FALLA'), name, '-', detail, flush=True)

from hp_fixture import Fake, PAYLOAD

fake = Fake()

def ctx_for(browser, device):
    ctx = browser.new_context(**{k: v for k, v in device.items() if k != 'default_browser_type'}, service_workers='allow', accept_downloads=True)
    ctx.add_init_script(f"if (!localStorage.getItem('{STORAGE_KEY}')) localStorage.setItem('{STORAGE_KEY}', {json.dumps(json.dumps(fake.session('A')))});")
    ctx.route(f'{SB}/**', fake.handle)
    return ctx

def bg(page, sel):
    return page.eval_on_selector(sel, 'e => getComputedStyle(e).backgroundColor')

def download(page, click):
    with page.expect_download(timeout=60000) as d: click()
    p = os.path.join(OUT, d.value.suggested_filename); d.value.save_as(p); return p

with sync_playwright() as pw:
    browser = pw.chromium.launch()
    a1, f1, b1 = fake.hp_audits[0]
    check('0.1 Datos de prueba: auditoría con las respuestas de la planilla = 6,27 Bueno', f1 == 6.27 and b1 == 'Bueno', f'{f1} {b1}')
    desk = ctx_for(browser, {'viewport': {'width': 1366, 'height': 900}})
    P = desk.new_page()
    P.goto(f'{BASE}/auditorias/{a1}'); P.wait_for_selector('.sheet-table', timeout=60000)
    body = P.inner_text('body')
    check('1.1 Planilla de resultado: 6,27 Bueno, oficial', '6,27' in body and 'Bueno' in body and 'Resultado oficial' in body)
    rows = P.eval_on_selector_all('.sheet-table tbody tr', 'rs => rs.map(r => [...r.cells].map(c => c.innerText.trim()))')
    exp = [('16', '24', '6,67'), ('63', '90', '7,00'), ('65', '84', '7,74'), ('4', '12', '3,33'), ('5', '9', '5,56'), ('22', '30', '7,33')]
    check('1.2 Requisitos con puntaje alcanzado, objetivo y evaluación iguales a la planilla', [tuple(r[1:4]) for r in rows[:6]] == exp, json.dumps(rows[:6], ensure_ascii=False)[:300])
    check('1.3 Fila de resultado final 175 / 249 / 6,27', rows[6][:4] == ['Resultado final', '175', '249', '6,27'], json.dumps(rows[6], ensure_ascii=False))
    colors = P.eval_on_selector_all('.sheet-table tbody td.eval', 'cs => cs.map(c => getComputedStyle(c).backgroundColor)')
    G, C, R = 'rgb(220, 250, 230)', 'rgb(254, 228, 226)', 'rgb(254, 240, 199)'  # tonos de pantalla (Bueno, Crítico, Regular)
    want = [G, G, G, C, R, G, G]
    check('1.4 Colores de evaluación por sección según la banda (Bueno, Crítico, Regular)', colors == want, json.dumps(colors))
    crit = P.inner_text('.criteria')
    check('1.5 Cuadro "Criterio de evaluación" con rangos y bandas', all(s in crit for s in ['8,01 - 10', 'Muy Bueno', '6,01 - 8', 'Bueno', '4,01 - 6', 'Regular', '0 - 4', 'Crítico']), crit.replace('\n', ' '))
    check('1.6 Recuadro del resultado con el color de la banda', bg(P, '.verdict') == 'rgb(220, 250, 230)', bg(P, '.verdict'))
    P.screenshot(path=f'{OUT}/1-planilla-escritorio.png', full_page=False)

    P.goto(f'{BASE}/'); P.wait_for_selector('text=Últimos resultados', timeout=60000); P.wait_for_timeout(800)
    chips = P.eval_on_selector_all('.sheet-table .band-chip', 'cs => cs.map(c => [c.innerText, getComputedStyle(c).backgroundColor])')
    labels = {c[0] for c in chips}
    check('2.1 Dashboard: últimos resultados con las cuatro bandas en color', {'Muy Bueno', 'Bueno', 'Regular', 'Crítico'} <= labels, json.dumps(chips, ensure_ascii=False)[:300])
    P.screenshot(path=f'{OUT}/2-dashboard.png', full_page=True)
    P.goto(f'{BASE}/auditorias'); P.wait_for_selector('text=AUD-2026-0012'); P.wait_for_timeout(500)
    P.screenshot(path=f'{OUT}/3-auditorias.png', full_page=False)

    # ---- comparación con la auditoría anterior (seguimiento de la misma empresa y plantilla)
    a4, f4, b4 = fake.hp_audits[3]
    P.goto(f'{BASE}/auditorias/{a4}'); P.wait_for_selector('.sheet-table', timeout=60000); P.wait_for_timeout(500)
    heads = P.eval_on_selector_all('.sheet-table thead th', 'hs => hs.map(h => h.innerText.trim())')
    tot = P.eval_on_selector_all('.sheet-table tr.total td', 'cs => cs.map(c => c.innerText.trim())')
    foot = P.inner_text('.sheet')
    check('6.1 Seguimiento: columnas Anterior y Variación con la auditoría previa de la misma empresa', 'Anterior' in heads and 'Variación' in heads and 'AUD-2026-0010' in foot, json.dumps(heads, ensure_ascii=False))
    check('6.2 Variación del resultado final (5,29 vs 6,27 → ▼ -0,98, en rojo)', tot[4] == '6,27' and '▼' in tot[5] and '-0,98' in tot[5] and P.eval_on_selector('.sheet-table tr.total .delta', 'e => getComputedStyle(e).color') == 'rgb(217, 45, 32)', json.dumps(tot, ensure_ascii=False))
    P.screenshot(path=f'{OUT}/7-comparacion.png', full_page=False)

    # ---- acta de reunión de cierre y firma en campo
    P.goto(f'{BASE}/auditorias/{a1}'); P.wait_for_selector('text=Reunión de cierre y firmas', timeout=60000)
    P.get_by_role('button', name='Expandir todo').click()
    P.wait_for_selector('.thumb .photo-num', timeout=30000)
    tags = P.eval_on_selector_all('.thumb .photo-num', 'xs => xs.map(x => x.innerText.trim())')
    check('3.8 Pantalla: la foto del requisito muestra su número ("Foto 1")', 'Foto 1' in tags, json.dumps(tags, ensure_ascii=False))
    P.locator('.thumb').first.scroll_into_view_if_needed(); P.screenshot(path=f'{OUT}/11-foto-numerada.png', full_page=False)
    P.get_by_label('Fecha de la reunión').fill('2026-05-07')
    P.get_by_label('Asistentes').fill('Juan Pérez (Contratista Patagonia SA, Jefe HSE); Auditor A')
    P.get_by_label('Acuerdos y compromisos').fill('Presentar el plan de acción en 15 días.')
    P.get_by_role('button', name='Guardar acta').click(); P.wait_for_timeout(400)
    for who, name in (('auditor', 'Auditor A'), ('contratista', 'Juan Pérez')):
        P.get_by_role('button', name='Agregar firma').click()
        if who == 'contratista':
            P.get_by_label('Nombre y apellido').fill(name)
            P.get_by_role('radio', name='Con observaciones').click(); P.get_by_label('Observaciones del firmante').fill('Pedimos 30 días para residuos.')
        pad = P.locator('canvas.sigpad'); bb = pad.bounding_box()
        P.mouse.move(bb['x'] + 30, bb['y'] + 120); P.mouse.down()
        for i in range(1, 12): P.mouse.move(bb['x'] + 30 + i * 25, bb['y'] + 120 - (40 if i % 2 else 0), steps=3)
        P.mouse.up()
        P.get_by_role('button', name='Guardar firma').click(); P.wait_for_timeout(500)
    figs = P.eval_on_selector_all('.sig figcaption strong', 'xs => xs.map(x => x.innerText)')
    check('7.1 Acta y dos firmas capturadas en el dispositivo (auditor y contratista)', figs == ['Auditor A', 'Juan Pérez'], json.dumps(figs, ensure_ascii=False))
    P.wait_for_timeout(2500)
    srv = [r for r in fake.tables['hse_audit_signatures'].values() if r.get('audit_id') == a1]
    check('7.2 Firmas sincronizadas como imagen PNG', len(srv) == 2 and all((r.get('signature_png') or '').startswith('data:image/png;base64,') for r in srv), f'{len(srv)} en el servidor')
    P.screenshot(path=f'{OUT}/8-firmas.png', full_page=False)

    P.goto(f'{BASE}/informes?audit={a1}'); P.wait_for_selector('text=Informe de una auditoría')
    P.wait_for_timeout(2500)
    pdf = download(P, lambda: P.get_by_role('button', name='Informe detallado (PDF)').click())
    txt = ' '.join(' '.join((pg.extract_text() or '').split()) for pg in PdfReader(pdf).pages)
    check('3.4 PDF: acta de reunión de cierre con firmas y conformidad', all(x in txt for x in ['Acta de reunión de cierre', 'Juan Pérez', 'Con observaciones', 'Presentar el plan de acción en 15 días']), '')
    check('3.0 PDF: carátula de RS Consultora, índice y sin historial de cambios', all(s in txt for s in ['RS CONSULTORA', 'Informe de auditoría', 'Índice', '1. Datos generales', 'Resumen ejecutivo', 'Página 3 de']) and 'Historial de cambios' not in txt, '')
    check('3.6 PDF: fotos numeradas, agrupadas por requisito y citadas en la lista de verificación', all(x in txt for x in ['Registro fotográfico', 'Foto 1 · Extintor sin tarjeta', 'Req. 1 ', '» Foto 1', 'Hallazgo HAL-2026-0001', 'Foto 2 (ver registro fotográfico)']), '')
    check('3.1 PDF: tabla de requisitos, resultado final y criterio de evaluación', all(s in txt for s in ['Requisitos del sistema de gestión', 'Puntaje alcanzado', 'RESULTADO FINAL 175 249 6,27', 'Criterio de evaluación', '8,01 - 10 Muy Bueno']), os.path.basename(pdf))
    xlsx = download(P, lambda: P.get_by_role('button', name='Descargar Excel').first.click())
    wbx = load_workbook(xlsx)
    ckr = [[c.value for c in row] for row in wbx['Checklist'].iter_rows(min_row=1)]
    check('3.7 Excel: columna "Fotos N.º" con los números de foto', ckr[0][6] == 'Fotos N.º' and any(r[6] == '1' for r in ckr[1:]), json.dumps([r[6] for r in ckr[1:4]], ensure_ascii=False))
    check('3.5 Excel: hoja "Acta y firmas"', 'Acta y firmas' in wbx.sheetnames and any(c.value == 'Juan Pérez' for c in wbx['Acta y firmas']['A']), ', '.join(wbx.sheetnames))
    ws = wbx['Secciones (oficial)']
    fills = [ws.cell(r, 4).fill.fgColor.rgb for r in range(2, 9)]
    check('3.2 Excel: evaluación coloreada por banda y fila de resultado final', fills == ['FF92D050', 'FF92D050', 'FF92D050', 'FFE26B0A', 'FFFFFF00', 'FF92D050', 'FF92D050'] and ws.cell(8, 1).value == 'Resultado final' and ws.cell(8, 1).fill.fgColor.rgb == 'FFD8E4BC', json.dumps(fills))
    check('3.3 Excel: encabezado con el color de la planilla', ws.cell(1, 1).fill.fgColor.rgb == 'FFFAC090', ws.cell(1, 1).fill.fgColor.rgb)
    # auditoría creada con las respuestas de la planilla importada
    P.goto(f'{BASE}/plantillas/{fake.hp_template}'); P.wait_for_selector('text=Crear auditoría con las respuestas de la planilla', timeout=60000)
    P.get_by_role('button', name='Crear auditoría con las respuestas de la planilla').click()
    P.get_by_label('Título').fill('CSMS a segundas partes – Contratista real')
    P.locator('.modal').get_by_role('button', name='Crear auditoría').click()
    P.wait_for_selector('.audit-bar', timeout=60000); P.wait_for_timeout(800)
    bar = P.locator('.audit-bar').inner_text().replace('\n', ' ')
    check('8.1 Auditoría creada con las 83 respuestas de la planilla y resultado parcial 6,27', '83/83' in bar and '6,27' in bar, bar)
    P.wait_for_timeout(3000)
    newa = [a for a in fake.tables['hse_audits'].values() if a.get('title') == 'CSMS a segundas partes – Contratista real']
    nresp = len([r for r in fake.tables['hse_audit_responses'].values() if newa and r.get('audit_id') == newa[0]['id']])
    check('8.2 Se sincroniza al servidor (auditoría en curso + 83 respuestas)', len(newa) == 1 and newa[0].get('status') == 'en_curso' and nresp == 83, f'{len(newa)} auditoría, {nresp} respuestas')
    # lo mismo desde Auditorías › Nueva auditoría (opción marcada por defecto)
    P.goto(f'{BASE}/auditorias'); P.wait_for_selector('text=Nueva auditoría', timeout=30000)
    P.get_by_role('button', name='Nueva auditoría').click()
    P.locator('.modal select').first.select_option(index=1)
    P.wait_for_selector('.sheet-opt', timeout=30000)
    P.get_by_label('Título').fill('CSMS desde Nueva auditoría')
    P.screenshot(path=f'{OUT}/9-nueva-con-planilla.png')
    P.locator('.modal').get_by_role('button', name='Crear', exact=True).click()
    P.wait_for_selector('.audit-bar', timeout=60000); P.wait_for_timeout(800)
    bar2 = P.locator('.audit-bar').inner_text().replace('\n', ' ')
    check('8.3 Nueva auditoría: opción "Cargar las respuestas de la planilla importada" (83/83, 6,27)', '83/83' in bar2 and '6,27' in bar2, bar2)
    desk.close()

    mob = ctx_for(browser, pw.devices['Pixel 7'])
    M = mob.new_page()
    M.goto(f'{BASE}/auditorias/{a1}'); M.wait_for_selector('.sheet-table', timeout=60000); M.wait_for_timeout(500)
    check('4.1 Celular: planilla sin desplazamiento horizontal de la página', M.evaluate('w => document.documentElement.scrollWidth <= w + 1', M.viewport_size['width']))
    M.screenshot(path=f'{OUT}/4-planilla-celular.png', full_page=False)
    M.goto(f'{BASE}/'); M.wait_for_selector('text=Últimos resultados', timeout=60000); M.wait_for_timeout(500)
    check('4.2 Celular: dashboard sin desplazamiento horizontal', M.evaluate('w => document.documentElement.scrollWidth <= w + 1', M.viewport_size['width']))
    M.screenshot(path=f'{OUT}/5-dashboard-celular.png', full_page=True)
    mob.close()

    out = browser.new_context(viewport={'width': 1200, 'height': 800})
    L = out.new_page(); L.goto(f'{BASE}/'); L.wait_for_selector('text=Iniciar sesión', timeout=30000); L.wait_for_timeout(500)
    check('5.1 Pantalla de ingreso con la tipografía nueva', L.evaluate("() => document.fonts.check('16px \"Archivo Variable\"')"))
    L.screenshot(path=f'{OUT}/6-ingreso.png')
    browser.close()

json.dump(results, open(f'{OUT}/resultados.json', 'w'), ensure_ascii=False, indent=2)
ok = sum(r['ok'] for r in results)
print(f'\n{ok}/{len(results)} pruebas OK')
sys.exit(0 if ok == len(results) else 1)
